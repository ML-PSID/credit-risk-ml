from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
IMAGE_DIR = BASE_DIR / "static" / "image-ml"
SOURCE = BASE_DIR / "rapport_section_rl.md"
OUTPUT = BASE_DIR / "rapport_section_rl.docx"

IMAGE_CAPTIONS = {
    "charts_rl_resultats.png": "Figure RL 1 - Matrice de confusion et métriques principales.",
    "charts_rl_roc_pr.png": "Figure RL 2 - Courbes ROC et précision-rappel.",
    "charts_rl_cv.png": "Figure RL 3 - Validation croisée stratifiée à 5 folds.",
    "charts_rl_learning.png": "Figure RL 4 - Courbe d'apprentissage.",
    "charts_rl_coefficients.png": "Figure RL 5 - Coefficients standardisés les plus importants.",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table) -> None:
    tbl_pr = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for name, value in {"top": "80", "bottom": "80", "start": "120", "end": "120"}.items():
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tbl_pr.append(margins)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(31, 77, 120)
    title.paragraph_format.space_after = Pt(10)

    for name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(46, 116, 181), 16, 8),
        ("Heading 2", 13, RGBColor(46, 116, 181), 12, 6),
        ("Heading 3", 12, RGBColor(31, 77, 120), 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(85, 85, 85)


def add_image(doc: Document, image_name: str) -> None:
    image_path = IMAGE_DIR / image_name
    if not image_path.exists():
        p = doc.add_paragraph()
        p.add_run(f"[Image manquante : {image_name}]").bold = True
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(5.9))
    add_caption(doc, IMAGE_CAPTIONS.get(image_name, image_name))


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    set_cell_margins(table)

    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = text
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
            if r_idx == 0:
                set_cell_shading(cell, "F2F4F7")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(31, 77, 120)

    doc.add_paragraph()


def add_markdown_paragraph(doc: Document, line: str) -> None:
    if line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(line[2:])
        return

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)

    if re.match(r"^[A-ZÉÈÀ].*  $", line):
        line = line.rstrip()

    parts = re.split(r"(`[^`]+`|\\*\\*[^*]+\\*\\*)", line)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(31, 77, 120)
        elif part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)


def build_docx() -> None:
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Régression Logistique pondérée")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Section à intégrer au rapport ML - Risque de défaut de paiement")
    run.italic = True
    run.font.color.rgb = RGBColor(85, 85, 85)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    idx = 0
    in_code = False
    table_rows: list[list[str]] = []

    while idx < len(lines):
        raw = lines[idx]
        line = raw.strip()

        if line.startswith("# 2. "):
            idx += 1
            continue

        if line.startswith("```"):
            in_code = not in_code
            idx += 1
            continue

        if in_code:
            if line:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                run = p.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(31, 77, 120)
            idx += 1
            continue

        if line.startswith("Image à insérer :"):
            match = re.search(r"`([^`]+\.png)`", line)
            if match:
                add_image(doc, match.group(1))
            else:
                add_markdown_paragraph(doc, line)
            idx += 1
            continue

        if line.startswith("## "):
            if line.startswith("## 2.7 "):
                doc.add_page_break()
            doc.add_heading(line[3:], level=1)
            idx += 1
            continue

        if line.startswith("|"):
            table_rows = []
            while idx < len(lines) and lines[idx].strip().startswith("|"):
                row_line = lines[idx].strip()
                if not re.match(r"^\|\s*:?-+", row_line):
                    cells = [cell.strip() for cell in row_line.strip("|").split("|")]
                    table_rows.append(cells)
                idx += 1
            add_table(doc, table_rows)
            continue

        if not line:
            idx += 1
            continue

        add_markdown_paragraph(doc, line)
        idx += 1

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
